#!/usr/bin/env python3
"""Genera el informe Word de la jornada del 2026-09-19.

    python3 docs/generar_informe.py

Lee docs/graficas/ (tools/graficas_navegacion.py), docs/sendero/
(tools/probar_sendero.py) y las fotos de docs/fotos_huerta/ si las hay.
Volver a correrlo después de copiar las fotos las incluye.
"""
import json
from datetime import date
from pathlib import Path

import cv2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
GRAF = DOCS / "graficas"
SEND = DOCS / "sendero"
FOTOS = DOCS / "fotos_huerta"
SALIDA = DOCS / "Informe_AgroBot_2026-09-19.docx"
VERDE = RGBColor(0x06, 0x7A, 0x35)
ANCHO_UTIL_CM = 16.0

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for lado in ("left_margin", "right_margin"):
    setattr(sec, lado, Cm(2.5))
sec.top_margin = sec.bottom_margin = Cm(2.2)

base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)
for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
    st = doc.styles[f"Heading {nivel}"]
    st.font.name, st.font.size, st.font.color.rgb = "Calibri", Pt(tam), VERDE

_figura = [0]
_tabla = [0]


def p(texto: str, negrita_inicio: str | None = None, estilo=None):
    par = doc.add_paragraph(style=estilo)
    if negrita_inicio:
        par.add_run(negrita_inicio).bold = True
    par.add_run(texto)
    par.paragraph_format.space_after = Pt(6)
    return par


def viñetas(items):
    for it in items:
        if isinstance(it, tuple):
            p(it[1], negrita_inicio=it[0], estilo="List Bullet")
        else:
            p(it, estilo="List Bullet")


def _sombrear(celda, color_hex):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc.append(shd)


def tabla(titulo, cabecera, filas, anchos_cm):
    _tabla[0] += 1
    cap = doc.add_paragraph()
    cap.add_run(f"Tabla {_tabla[0]}. ").bold = True
    cap.add_run(titulo)
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(cabecera):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(txt))
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(c, "1A9850")
    for k, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, txt in enumerate(fila):
            celdas[i].text = ""
            celdas[i].paragraphs[0].add_run("—" if txt is None else str(txt)).font.size = Pt(9)
            if k % 2:
                _sombrear(celdas[i], "EEF6F0")
    for fila in t.rows:
        for i, ancho in enumerate(anchos_cm):
            fila.cells[i].width = Cm(ancho)
    doc.add_paragraph()


def figura(ruta: Path, pie: str, ancho_cm=ANCHO_UTIL_CM):
    if not ruta.exists():
        p(f"[Falta la imagen {ruta.name}]")
        return
    if ruta.suffix.lower() not in (".png", ".jpg", ".jpeg"):
        tmp = ruta.with_suffix(".conv.jpg")
        cv2.imwrite(str(tmp), cv2.imread(str(ruta)))
        ruta = tmp
    _figura[0] += 1
    doc.add_picture(str(ruta), width=Cm(ancho_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Figura {_figura[0]}. ")
    r.bold = True
    r.font.size = Pt(9)
    cap.add_run(pie).font.size = Pt(9)


def fotos_huerta():
    fotos = sorted(f for f in FOTOS.glob("*") if f.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp"))
    return [f for f in fotos if not f.name.endswith(".conv.jpg")]


metricas = json.loads((GRAF / "metricas.json").read_text())
pruebas_sendero = json.loads((SEND / "resultados.json").read_text())

# ---------------------------------------------------------------- portada
titulo = doc.add_paragraph()
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = titulo.add_run("AgroBot — Informe de pruebas y mejoras")
r.bold, r.font.size, r.font.color.rgb = True, Pt(22), VERDE
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Navegación, LIDAR, manejo, detección de senderos y dashboard de entrenamiento\n"
            "Jornada del 19 de septiembre de 2026").italic = True
p(f"Generado el {date.today().isoformat()} con docs/generar_informe.py a partir de los registros "
  "del robot (~/.agrobot/navigation_logs), las gráficas de tools/graficas_navegacion.py y las pruebas "
  "de tools/probar_sendero.py.")

doc.add_heading("Resumen", level=1)
s4, s5 = metricas["sesiones"][3], metricas["sesiones"][4]
viñetas([
    ("Navegación: ", f"en las dos pruebas largas la pose final quedó a {s4['deriva_m']} m y "
     f"{s5['deriva_m']} m del inicio ({s4['deriva_pct']} % y {s5['deriva_pct']} % de lo recorrido). "
     "Se ajustó todo el SLAM, el EKF y Nav2 para reducirlo."),
    ("LIDAR: ", "tres fallos encadenados explicaban que adelante/atrás no funcionaran: el USB se "
     "reenumeraba y el driver quedaba colgado, el freno veía el propio chasis y el LIDAR estaba "
     "girado 180° respecto de lo que suponía el software. Los tres están corregidos."),
    ("Manejo: ", "modos Interior y Exterior; la maniobra de giro ahora desplaza un cuarto de lo que "
     "desplazaba y ya no pisa las plantas."),
    ("Senderos: ", f"detector nuevo por cámara pensado para la huerta real (tierra oscura bordeada de "
     f"paja). En {pruebas_sendero['resumen']['pruebas']} escenas de prueba no tuvo fallos y midió el "
     f"centro del sendero con {pruebas_sendero['resumen']['error_lateral_medio_cm']} cm de error medio."),
    ("Entrenamiento: ", "dashboard aparte con mapa en vivo, corrección de pose, dibujo de surcos, "
     "puntos de escaneo y registro de correcciones."),
])

# ---------------------------------------------------------------- huerta
doc.add_heading("1. Zona de pruebas exterior", level=1)
p("La huerta de pruebas es un sendero en bucle de tierra desnuda que rodea un cantero central con "
  "maíz, brócoli y acelga. Lo que delimita el sendero a cada lado es, sobre todo, paja seca (mulch) "
  "de color claro, con plantas todavía pequeñas y escasas. En un extremo el sendero se abre en una "
  "zona sin bordes claros, donde se hace el giro del bucle.")
fotos = fotos_huerta()
if fotos:
    for f in fotos:
        figura(f, f"Huerta de pruebas ({f.stem}).", ancho_cm=12)
else:
    p("[Las fotos de la huerta van aquí: copiarlas a docs/fotos_huerta/ y volver a ejecutar "
      "docs/generar_informe.py.]").runs[0].italic = True
tabla("Características de la huerta que condicionan la navegación.",
      ["Rasgo", "Consecuencia para el robot"],
      [["Plantas pequeñas (< 30 cm)", "El LIDAR, montado a ~30 cm, pasa por encima: no sirve para seguir hileras."],
       ["Paja clara a los lados del camino", "Es el borde más visible: el detector de la cámara la usa."],
       ["Tierra oscura en el sendero", "Buen contraste con la paja; hay pajas sueltas que hay que ignorar."],
       ["Bucle con curvas", "El detector busca el rumbo del sendero hasta ±40°."],
       ["Zona abierta en un extremo", "Sin bordes la cámara se abstiene; manda el LIDAR y el mapa."],
       ["Cercas, paredes y estructuras", "Buenas referencias para el SLAM con LIDAR."]],
      [5.5, 10.5])

# ---------------------------------------------------------------- navegación
doc.add_heading("2. Pruebas de navegación y desplazamiento", level=1)
p("Las sesiones de calibración registran la pose del EKF, las velocidades y las distancias del "
  "LIDAR. En estas pruebas el robot vuelve al punto de partida, así que la distancia final al inicio "
  "(deriva) mide directamente el error acumulado de la localización.")
filas = []
for i, s in enumerate(metricas["sesiones"], 1):
    filas.append([f"#{i}", s["nombre"], s["poses"], s["duracion_s"], s["distancia_m"], s["deriva_m"],
                  None if s["deriva_pct"] is None else f"{s['deriva_pct']} %", s["lidar_frontal_cm"]])
b = metricas.get("bug_imu")
if b:
    filas.append(["bug IMU", "reloj en 1970", b["poses"], b["duracion_s"], b["distancia_m"], b["deriva_m"],
                  f"{b['deriva_pct']} %", b["lidar_frontal_cm"]])
tabla("Métricas por sesión de navegación.",
      ["Sesión", "Fecha/hora", "Poses", "Duración (s)", "Recorrido (m)", "Deriva (m)", "Deriva / recorrido",
       "LIDAR frontal medio (cm)"], filas, [1.5, 2.8, 1.3, 1.7, 2.0, 1.7, 2.3, 2.7])
p("Las sesiones #1 a #3 son pruebas cortas con el robot casi quieto: la deriva es de milímetros. "
  "En las sesiones largas (#4 y #5) la deriva crece con la distancia. La fila 'bug IMU' es una sesión "
  "antigua con la aceleración del MPU6050 fusionada en el EKF, donde la pose se iba a más de un "
  "kilómetro: es el antes y el después de haber quitado esa fuente.")
for nombre, pie in (("01_trayectorias.png", "Trayectorias XY de cada sesión: inicio (verde) y fin (rojo)."),
                    ("02_desplazamiento.png", "Distancia al punto de inicio a lo largo del tiempo."),
                    ("06_resumen.png", "Resumen por sesión: recorrido, deriva, duración y LIDAR frontal."),
                    ("04_rumbo.png", "Rumbo estimado a lo largo del tiempo."),
                    ("03_lidar_distribucion.png", "Distribución de distancias del LIDAR por sector."),
                    ("05_lidar_frontal.png", "Distancia frontal del LIDAR en las sesiones largas."),
                    ("07_antes_despues.png", "Antes y después de quitar la aceleración de la IMU del EKF.")):
    figura(GRAF / nombre, pie)

doc.add_heading("2.1 Ajustes aplicados a la localización", level=2)
tabla("Parámetros de SLAM, EKF y Nav2 cambiados tras analizar las pruebas.",
      ["Componente", "Antes", "Después", "Por qué"],
      [["slam_toolbox: distancia mínima entre nodos", "0,30 m", "0,12 m", "Grafo más denso, menos deriva entre nodos"],
       ["slam_toolbox: giro mínimo entre nodos", "0,30 rad", "0,15 rad", "Los giros en sitio generan nodos"],
       ["slam_toolbox: cierre de bucles", "desactivado", "activo (4 m, cadena 3)", "Corrige la deriva al volver por el bucle"],
       ["slam_toolbox: alcance del láser", "12 m", "8 m", "El C1 es ruidoso más allá"],
       ["Entrada del SLAM", "/scan crudo", "/scan_slam (rango + mediana)", "Quita picos de ruido"],
       ["EKF: ruido de proceso", "por defecto", "matriz explícita", "Menos confianza en rf2o (vx), más en el giroscopio"],
       ["Nav2: controlador", "DWB", "RotationShim + DWB", "Gira primero hacia la ruta y luego avanza"],
       ["Nav2: costmap local", "4 × 4 m", "5 × 5 m", "Ve antes los obstáculos"],
       ["Patrulla", "arranca en seco", "3 s de estabilización", "El SLAM se asienta antes de moverse"]],
      [4.5, 2.5, 3.5, 5.5])

# ---------------------------------------------------------------- fallos del día
doc.add_heading("3. Problemas encontrados y soluciones", level=1)
tabla("Fallos de la jornada, causa y corrección.",
      ["Síntoma", "Causa", "Solución"],
      [["Adelante/atrás no encendían los motores; los giros sí",
        "Sin datos del LIDAR el freno corta la velocidad lineal (nunca el giro)",
        "Ver las tres filas siguientes"],
       ["El LIDAR dejaba de leer", "El CP2102N se cae del USB (4 veces en 15 min) y vuelve como otro ttyUSB",
        "Puerto por número de serie (/dev/serial/by-id)"],
       ["Tras reconectar seguía sin leer", "El driver sllidar queda colgado sin morir: el respawn no actúa",
        "Vigía en el backend: sin scan 4 s con el USB presente, relanza el driver"],
       ["Frenaba sin obstáculo delante", "Reflejos del propio chasis (5–15 cm) contados como obstáculo",
        "Se ignoran los puntos dentro de la huella del robot"],
       ["Frenaba por una pared al costado", "Cono frontal de 45° incluía la pared paralela",
        "Adelante/atrás solo cuentan lo que está en el camino que barre el robot (+5 cm)"],
       ["Delante y detrás invertidos, y también izquierda y derecha",
        "Con la flecha del C1 al frente, el 0° del driver apunta hacia atrás",
        "Giro de 180° en el TF base_link→laser y en el backend"],
       ["El dashboard de entrenamiento casi no movía el robot",
        "Mandaba 'parar' al soltar el botón y cortaba el pulso (y la patrulla)",
        "Un toque = un pulso, como el dashboard principal"],
       ["Los botones Interior/Exterior no respondían", "La función quedaba fuera del alcance del onclick",
        "Eventos por código; ambos dashboards se sincronizan"],
       ["La maniobra de giro pisaba plantas", "Arcos de 0,6 s al 50 % de velocidad",
        "Modo Exterior: arcos de 0,3 s al 25 % (un cuarto del desplazamiento)"]],
      [4.8, 5.6, 5.6])
p("Queda un problema físico: el adaptador USB del LIDAR se desconecta. El vigía lo recupera solo, "
  "pero conviene revisar el cable, el hub y la alimentación, sobre todo si coincide con arrancar "
  "los motores.")

# ---------------------------------------------------------------- manejo
doc.add_heading("4. Modos de manejo", level=1)
tabla("Parámetros de cada modo (pulsos del manejo manual).",
      ["Parámetro", "Interior", "Exterior"],
      [["Velocidad", "60 % del máximo", "100 % del máximo"],
       ["Pulso adelante/atrás", "0,8 s", "1,0 s"],
       ["Pulso de giro en su eje", "0,4 s", "0,35 s"],
       ["Maniobra de giro", "sin avance (gira en su eje)", "arcos de 0,3 s al 25 % de velocidad"],
       ["Uso", "suelo liso, sitios estrechos", "tierra blanda, donde girar en el sitio entierra las ruedas"]],
      [5, 5.5, 5.5])

# ---------------------------------------------------------------- senderos
doc.add_heading("5. Detección de senderos por cámara", level=1)
p("La versión anterior seguía el centro de lo que no era verde en la parte baja de la imagen. Con "
  "plantas pequeñas casi toda la imagen es tierra y ese centro caía siempre en medio: no medía nada. "
  "Además incluía el borde inferior, donde se ve el propio chasis.")
doc.add_heading("5.1 Cómo funciona ahora", level=2)
viñetas([
    ("Bordes: ", "plantas (índice Excess Green con umbral de Otsu, más el rango de verde configurable) "
     "y paja (píxeles más claros que el suelo de esa misma imagen y de tono amarillento). El umbral se "
     "adapta a cada fotograma, así aguanta sol y sombra."),
    ("Vista cenital: ", "los bordes se proyectan al suelo, en metros, con la altura, inclinación y campo "
     "de la cámara. En el suelo los bordes son paralelos y las pajas sueltas no forman línea."),
    ("Rumbo: ", "se prueba cada ángulo entre −40° y +40° y se queda el que deja los bordes más alineados."),
    ("Sendero: ", "el hueco libre entre bordes que contiene la posición anterior; da centro y ancho reales. "
     "Con un solo borde visible usa el ancho medido antes."),
    ("Corrección: ", "el desvío se mide 0,6 m por delante del robot, en centímetros y con el mismo signo que "
     "el del LIDAR, y se compara con la tolerancia configurada."),
    ("Seguridad: ", "si la cámara está girada o escaneando, la visión se abstiene."),
])
figura(SEND / "deteccion_sintetica.jpg",
       "Detector sobre escenas sintéticas: arriba, robot 10 cm fuera del centro; abajo, sendero en curva. "
       "Izquierda: cámara con paja (amarillo) y plantas (verde). Derecha: vista cenital con los bordes "
       "(naranja), el centro del sendero (magenta) y el eje del robot (blanco).")
figura(SEND / "escena_y_mascara_paja.jpg",
       "Escenas de prueba (arriba) y máscara de paja detectada (abajo).")
figura(SEND / "deteccion_camara_real.jpg",
       "Cámara real del robot en un pasillo con plantas a un lado: detecta las plantas y se abstiene "
       "porque solo hay un borde visible.")

doc.add_heading("5.2 Resultados de las pruebas", level=2)
res = pruebas_sendero["resumen"]
p(f"{res['pruebas']} escenas ({len(pruebas_sendero['casos']) // 3} casos × 3 variantes de textura), "
  f"sendero de {res['ancho_real_cm']:.0f} cm con franjas de paja de 35 cm, pajas sueltas en el camino, "
  f"plantas escasas y sombra. Fallos: {res['fallos']}. Error medio del centro en tramos rectos: "
  f"{res['error_lateral_medio_cm']} cm.")
filas = [[c["caso"], c["lateral_real_cm"], c["lateral_medido_cm"], c["ancho_medido_cm"],
          c["rumbo_real_deg"], c["rumbo_medido_deg"], c["esperado"], c["decision"]]
         for c in pruebas_sendero["casos"] if c["semilla"] == 1]
tabla("Detector de senderos: valor real frente a medido (variante 1 de cada caso). Signos: + = a la izquierda.",
      ["Caso", "Centro real (cm)", "Centro medido (cm)", "Ancho medido (cm)", "Rumbo real (°)",
       "Rumbo medido (°)", "Giro esperado", "Decisión"], filas, [3.6, 1.6, 1.7, 1.6, 1.6, 1.7, 2.1, 2.1])
p("Donde el giro esperado es 'seguir' pero el detector corrige (o al revés), el desvío real está a "
  "menos de 3 cm de la tolerancia de 10 cm: cualquiera de las dos decisiones es válida.")

doc.add_heading("5.3 Ajuste en campo", level=2)
viñetas([
    "En el dashboard de entrenamiento, 'Vista del sendero' muestra en vivo lo que detecta la cámara.",
    "Si la paja no se marca en amarillo, o el camino se llena de amarillo, revisar la inclinación "
    "(odometria_webcam.camara_inclinacion_deg) y la altura de la cámara en Parámetros.",
    "Si el camino mismo está cubierto de paja, apagar 'La paja seca marca el borde'.",
    "Para repetir las pruebas: python3 tools/probar_sendero.py (sale con error si alguna decisión falla).",
])
tabla("Parámetros del seguidor de senderos.",
      ["Parámetro", "Valor", "Efecto"],
      [["sendero_tolerancia_cm", "10", "Desvío que se tolera sin corregir"],
       ["sendero_vision_peso", "0,5", "0 = solo LIDAR, 1 = solo cámara"],
       ["sendero_paja", "activo", "La paja cuenta como borde del sendero"],
       ["sendero_verde_h_min / h_max / s_min", "30 / 90 / 60", "Rango de verde de las plantas"],
       ["Rango de análisis", "0,3–1,8 m por delante", "Más lejos la resolución no alcanza"],
       ["Búsqueda de rumbo", "±40° en pasos de 2°", "Permite seguir las curvas del bucle"],
       ["Punto de mirada", "0,6 m", "Pesa a la vez desplazamiento y rumbo"]],
      [5.5, 3.5, 7])

# ---------------------------------------------------------------- dashboard
doc.add_heading("6. Dashboard de entrenamiento", level=1)
p("Se abre desde el enlace 'Entrenamiento' de la barra lateral (http://192.168.12.1:8080/training.html) "
  "y corre en paralelo al dashboard principal.")
viñetas(["Mapa del SLAM en vivo con la pose del robot.",
         "Corregir la pose con un clic (publica /initialpose).",
         "Dibujar surcos y colocar puntos de escaneo 🌱; generar la ruta de patrulla a partir de ellos.",
         "Manejo manual por pulsos con modo Interior/Exterior y maniobra activable.",
         "Configuración de conos, márgenes y velocidades.",
         "Vista del sendero en vivo.",
         "Registro de cada corrección del operador para el aprendizaje."])
tabla("Endpoints nuevos del backend.",
      ["Endpoint", "Función"],
      [["POST /api/training/correct_pose", "Corrige la pose en el SLAM"],
       ["POST /api/training/save_furrows", "Guarda los surcos dibujados"],
       ["POST /api/training/save_scan_points", "Guarda los puntos de escaneo"],
       ["POST /api/training/generate_route", "Genera waypoints en zigzag y los carga en la patrulla"],
       ["GET /api/training/data", "Surcos, puntos y correcciones guardados"],
       ["POST /api/training/log_correction", "Registra una corrección de manejo"],
       ["GET/POST /api/drive/mode", "Modo de manejo Interior/Exterior"],
       ["GET /api/sendero/vista.jpg", "Imagen de depuración del detector de senderos"]],
      [6.5, 9.5])

# ---------------------------------------------------------------- simulador
doc.add_heading("7. Propuesta: simulador y aprendizaje", level=1)
p("Se puede entrenar al robot en un simulador en otra PC y traer a la Jetson solo el modelo "
  "entrenado. Propuesta, pendiente de decidir:")
tabla("Arquitectura propuesta.",
      ["Pieza", "Dónde", "Qué hace"],
      [["Gazebo Harmonic", "PC con GPU", "Simulador nativo de ROS 2 Jazzy: robot skid-steer, RPLIDAR C1 y cámara "
        "a la misma altura e inclinación"],
       ["Réplica de la huerta", "PC", "El bucle con texturas de tierra y paja sacadas de las fotos"],
       ["Datos", "Jetson → PC", "CSV/JSONL que ya graba el robot (poses, LIDAR, correcciones del operador)"],
       ["Entrenamiento", "PC", "Imitación a partir de las correcciones reales y del manejo en el simulador"],
       ["Modelo", "PC → Jetson", "Archivo ONNX que la Jetson ejecuta con TensorRT"]],
      [3.5, 3, 9.5])
p("El riesgo principal es que la cámara simulada no se parezca a la real. Se mitiga variando al azar "
  "la luz y las texturas, y usando como entrada la misma vista cenital de bordes del detector, que "
  "se parece mucho más entre simulación y realidad que la imagen cruda.")

# ---------------------------------------------------------------- pendientes
doc.add_heading("8. Próximos pasos", level=1)
viñetas(["Probar el detector de senderos en la huerta y ajustar la geometría de la cámara con la vista en vivo.",
         "Recorrer el bucle una vez grabando la ruta para tener los waypoints sobre el mapa.",
         "Revisar la conexión física del LIDAR (cable, hub y alimentación).",
         "Medir el ancho real del robot con ruedas y ajustarlo en Parámetros.",
         "Decidir si se construye el simulador."])

doc.save(SALIDA)
print(f"{SALIDA} — {_figura[0]} figuras, {_tabla[0]} tablas, {len(fotos)} fotos de la huerta")
