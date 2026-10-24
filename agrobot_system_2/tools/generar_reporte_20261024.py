"""Informe unificado (Word) del 2026-10-24: auditoría de datos, mapas nuevos, entrenamiento CEM, percepción de coles, digitalización y
comparación real/simulación. Lee lo que haya en disco (informe.json del entrenamiento, métricas YOLO, validación de la digitalización)
y marca como PENDIENTE lo que aún no existe. Uso: python3 tools/generar_reporte_20261024.py [--intensivo <carpeta>]"""
import argparse, glob, json, os, sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parents[1]
OUT = RAIZ / "docs" / "REPORTE_MEJORAS_20261024"
FIG = OUT / "figs"
HOME = Path.home()


def js(p):
    try:
        return json.loads(Path(p).read_text())
    except Exception:
        return None


def tabla(doc, filas, cab, anchos=None):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(cab):
        t.rows[0].cells[i].text = str(c)
    for f in filas:
        r = t.add_row().cells
        for i, v in enumerate(f):
            r[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
    return t


def figura(doc, ruta, ancho=16, pie=None):
    if Path(ruta).exists():
        doc.add_picture(str(ruta), width=Cm(ancho))
        if pie:
            p = doc.add_paragraph(pie)
            p.runs[0].italic = True
            p.runs[0].font.size = Pt(8.5)
    else:
        doc.add_paragraph(f"[PENDIENTE: falta {Path(ruta).name}]")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--intensivo", default=None)
    a = ap.parse_args()
    FIG.mkdir(parents=True, exist_ok=True)
    carpetas = sorted(glob.glob(str(HOME / ".agrobot" / "aprendizaje" / "intensivo_20261024_*")))
    carp = Path(a.intensivo) if a.intensivo else (Path(carpetas[-1]) if carpetas else None)
    inf = js(carp / "informe.json") if carp else None
    his = js(carp / "historial.json") if carp else None
    dig = js(RAIZ / "docs" / "digitalizacion" / "resultados.json")
    yolo = js(RAIZ / "simulador_pasivo" / "percepcion_sintetica" / "modelos" / "col_plagas_v2_metricas.json")
    col_prev = js(RAIZ / "simulador_pasivo" / "percepcion_sintetica" / "modelos" / "col_10h_metricas.json")
    slam = js(RAIZ / "simulador_pasivo" / "resultados_slam_kaggle_20261023.json")

    doc = Document()
    doc.add_heading("AgroBot — Informe unificado de mejoras (2026-10-24)", 0)
    doc.add_paragraph("Auditoría de lo recopilado en las últimas pruebas, entrenamiento en simulación calibrada (interior y exterior, "
                      "con mapas nuevos), detector de plagas de col, digitalización del mapa del robot y comparación con fotos/videos reales. "
                      "Todo lo que NO se probó en el robot real está marcado como tal.")

    # ---------------------------------------------------------------- 1 datos auditados
    doc.add_heading("1. Qué dicen los datos de las últimas pruebas", 1)
    tabla(doc, [
        ["Robot parado ~90 % de la prueba del 22-oct", "89 % de filas con v≈0 (explorador no activo)", "FEEDBACK_PRUEBA_20261022.md"],
        ["Falsos positivos de plagas sobre fondos", "82/82 capturas finales sin planta; ~20 alarmas/min parado", "idem"],
        ["Imagen oscura y borrosa", "brillo medio 65/255; 89 % de fotos con nitidez < 60", "idem"],
        ["SLAM interior (22-oct)", "cierre 25-36 cm en 23 m; error ≈ 1-1,5 % del recorrido", "memoria agrobot-slam-interior-20261022"],
        ["Corrector ML del SLAM", "empeora al ICP crudo (ATE 0,37 vs 0,11 m en la sesión retenida)", "resultados_slam_kaggle_20261023.json"],
        ["Simulación calibrada (ronda 2, 23-oct)", "78,2 % éxito / 0,29 contactos vs HOY 56,4 % / 0,78; falencias: eslalon 37,5 %, rocas 42 %", "intensivo_20261023_1015"],
        ["Detector de col previo (col_10h)", "mAP50 0,59; dano_hoja 0,17; sin clase pulgón", "col_10h_metricas.json"],
    ], ["Hallazgo", "Dato", "Fuente"])
    if slam:
        doc.add_paragraph("SLAM: ATE (m) por sesión — base (ICP crudo) / corrector existente / sim+real:")
        tabla(doc, [[k, v["muestras"], v["recorrido_m"], v["base"]["ATE_m"], v["existente"]["ATE_m"], v["sim+real"]["ATE_m"]]
                    for k, v in slam["sesiones"].items()], ["sesión", "muestras", "m", "ATE base", "ATE existente", "ATE sim+real"])
        doc.add_paragraph("Conclusión: el ICP crudo es lo mejor medido; el corrector de aprendizaje NO debe activarse (slam.correccion_ml_activa sigue apagado).")

    # ---------------------------------------------------------------- 2 mapas nuevos
    doc.add_heading("2. Mapas de interior nuevos en el simulador", 1)
    doc.add_paragraph("simulador_pasivo/mapas_interiores.py: 4 laberintos/pasillos con esquinas de 90° y retorno en U, 2 laboratorios (bancos; mesas y sillas), "
                      "sala con puerta de 90 cm, puerta descentrada de 70 cm y almacén con estanterías. Se exportan a mundos/interior_*.json para el simulador "
                      "visual (suelo de baldosa, muros enyesados, muebles macizos, luminarias).")
    figura(doc, RAIZ / "docs" / "mapas_interiores.png", 16, "Mapas nuevos (verde = salida, estrella = meta).")
    doc.add_paragraph("Hallazgo clave: con la configuración de HOY el robot simulado no resuelve casi ninguna esquina de 90° (éxito 0/4 en pasillo_L, laberintos y laboratorio). "
                      "Las curvas suaves de 60° sí (4/4). Con 1 pulso de giro el robot avanza en arco y se sale por la pared exterior de la esquina; el "
                      "explorador real solo gira 1 pulso al bloquearse. Por eso el entrenamiento incluye las maniobras (giro de varios pulsos, retroceso, histéresis, pausa, escape).")

    # ---------------------------------------------------------------- 3 entrenamiento
    doc.add_heading("3. Entrenamiento Monte Carlo / CEM (interior + exterior)", 1)
    if inf:
        doc.add_paragraph(f"Carpeta {carp.name}: {inf['generaciones']} generaciones en {inf['duracion_min']} min. Validación en semillas nunca vistas.")
        h = inf["historial"]
        fig, ax = plt.subplots(figsize=(7, 3.2))
        ax.plot([x["gen"] for x in h], [x["mejor"] for x in h], label="mejor"); ax.plot([x["gen"] for x in h], [x["media"] for x in h], label="media")
        ax.set_xlabel("generación"); ax.set_ylabel("coste (menor = mejor)"); ax.legend(); ax.grid(alpha=.3); fig.tight_layout()
        fig.savefig(FIG / "cem_historial.png", dpi=130); plt.close(fig)
        figura(doc, FIG / "cem_historial.png", 13, "Evolución del coste durante la búsqueda.")
        filas = [[f["nombre"], f["todo"]["exito_pct"], f["todo"]["contactos_por_ep"], f["interior"]["exito_pct"], f["exterior"]["exito_pct"]]
                 for f in inf["validacion"]]
        tabla(doc, filas, ["candidato", "éxito % (todo)", "contactos/ep", "éxito % interior", "éxito % exterior"])
        nom = inf["mejor"]["nombre"]
        vm = next(f for f in inf["validacion"] if f["nombre"] == nom)
        vh = next(f for f in inf["validacion"] if f["nombre"] == "HOY")
        esc = sorted(vm["por_escenario"])
        fig, ax = plt.subplots(figsize=(10, 4.2))
        x = np.arange(len(esc))
        ax.bar(x - .2, [vh["por_escenario"][k]["exito_pct"] for k in esc], .4, label="HOY")
        ax.bar(x + .2, [vm["por_escenario"][k]["exito_pct"] for k in esc], .4, label=f"mejor ({nom})")
        ax.set_xticks(x); ax.set_xticklabels([k.split("/")[1] for k in esc], rotation=60, ha="right", fontsize=7); ax.set_ylabel("éxito %"); ax.legend(); fig.tight_layout()
        fig.savefig(FIG / "cem_por_escenario.png", dpi=130); plt.close(fig)
        figura(doc, FIG / "cem_por_escenario.png", 16, "Éxito por escenario en semillas nuevas: configuración de hoy vs mejor candidato.")
        doc.add_paragraph("Estrés (hardware flojo: más masa, más arrastre, menos agarre): " +
                          "; ".join(f"{k}: éxito {v['todo']['exito_pct']} %, contactos/ep {v['todo']['contactos_por_ep']}" for k, v in inf["estres"].items()))
        doc.add_paragraph("Parámetros del mejor candidato (NO desplegados; las maniobras man_* solo existen en el simulador hasta portarlas al explorador):")
        doc.add_paragraph(json.dumps(inf["mejor"]["params"], ensure_ascii=False))
    elif his:
        h = his["historial"]
        doc.add_paragraph(f"EN CURSO: {len(h)} generaciones; mejor coste {h[-1]['mejor']}, HOY {h[0].get('hoy')}.")
    else:
        doc.add_paragraph("[PENDIENTE: el entrenamiento no ha terminado]")

    # ---------------------------------------------------------------- 4 percepción
    doc.add_heading("4. Detección de plagas en coles", 1)
    doc.add_paragraph("Datos nuevos: aphid-generalization (626 + 260 fotos de pulgón con máscaras → cajas), AgroPest-12, y los 272 fondos sin planta como negativos. "
                      "El detector v2 añade la clase pulgón (la plaga real de las coles del operador: Brevicoryne brassicae).")
    if yolo:
        filas = [[k, (col_prev or {}).get("por_clase_mAP50", {}).get(k, "—"), v] for k, v in yolo["por_clase_mAP50"].items()]
        tabla(doc, filas, ["clase", "mAP50 previo (col_10h)", "mAP50 v2"])
        doc.add_paragraph(f"v2: mAP50 {yolo['mAP50']}, precisión {yolo['precision']}, recall {yolo['recall']}. Falsos positivos: {json.dumps(yolo['falsos_positivos'], ensure_ascii=False)}")
        doc.add_paragraph(yolo["NO_DEMOSTRADO"])
    else:
        doc.add_paragraph("[PENDIENTE: el entrenamiento YOLO no ha terminado]")

    # ---------------------------------------------------------------- 5 digitalización
    doc.add_heading("5. Digitalización del mapa del robot", 1)
    doc.add_paragraph("jetson/backend/digitalizar_mapa.py convierte el mapa de ocupación (PGM) en un plano vectorial: muros como segmentos rectos (Hough + fusión + "
                      "ajuste a los ejes del edificio), esquinas, espacios y área. Salidas SVG a escala, DXF (AutoCAD), JSON y PNG. API: POST /api/mapas/digitalizar; "
                      "botón «Digitalizar» en la lista de mapas del dashboard.")
    if dig:
        niveles = ["sin_deriva", "deriva_1pct", "deriva_3pct"]
        filas = []
        for n in niveles:
            l = [f for f in dig if f["nivel"] == n]
            filas.append([n, round(np.mean([f["recall_pct"] for f in l]), 1), round(np.mean([f["precision_pct"] for f in l]), 1),
                          round(np.mean([f["error_medio_cm"] for f in l if f["error_medio_cm"] is not None]), 1)])
        tabla(doc, filas, ["deriva de la pose", "recall % (≤10 cm)", "precisión %", "error medio cm"])
        doc.add_paragraph("Verdad conocida: los 10 mapas simulados, recorridos con el LIDAR calibrado y una pose con deriva (el SLAM real mide ~1-1,5 %). "
                          "Sin deriva el plano coincide al 93-95 %; con deriva la calidad cae: el cuello de botella es el SLAM, no la vectorización.")
    for n, cap in (("dig_real", "Mapa real guardado (mision_20261023_1618) digitalizado: 24 muros, 17 m."),):
        figura(doc, RAIZ / "docs" / "digitalizacion" / "plano_real_1618.png", 12, cap)

    # ---------------------------------------------------------------- 6 datasets
    doc.add_heading("6. Datasets descargados", 1)
    med = RAIZ / "datos_plantas_medicinales"
    filas = []
    for d in sorted(med.glob("*")):
        if d.is_dir():
            n = sum(1 for _ in d.rglob("*") if _.suffix.lower() in (".jpg", ".jpeg", ".png"))
            filas.append([d.name, n, "kaggle trientran/oriental-medicinal-herb-images (CC-BY-NC-SA-4.0)" if "oriental" in d.name
                          else "kaggle warcoder/indian-medicinal-plant-image-dataset"])
    tabla(doc, filas, ["dataset de plantas medicinales", "imágenes", "origen"])
    tabla(doc, [["aphid_generalization", "kaggle chahineeee/aphid-generalization (pulgón con máscaras)"],
                ["agropest12", "kaggle rupankarmajumdar/crop-pests-dataset (12 plagas con cajas)"]], ["plagas (extra)", "origen"])
    doc.add_paragraph("Son fotos de otras plantas y cámaras: sirven para preentrenar y medir, no sustituyen fotos de la planta del operador. Verificar licencia antes de redistribuir.")

    # ---------------------------------------------------------------- 7 comparación real vs sim
    doc.add_heading("7. Contraste con fotos y videos reales", 1)
    cmp_dir = RAIZ / "docs" / "comparacion_real_sim"
    for f, cap in (("real_robot_0.jpg", "Cámara del robot, prueba real del 22-oct."), ("real_video_8_09_Pro.jpg", "Video real del operador (suelo)."),
                   ("sim_interior_laboratorio_mesas_sillas.png", "Simulación: laboratorio con mesas y sillas (suelo de baldosa nuevo)."),
                   ("sim_interior_laberinto_serpiente.png", "Simulación: laberinto en serpiente.")):
        figura(doc, cmp_dir / f, 9, cap)
    doc.add_paragraph("Comparación honesta: el LIDAR simulado usa el ruido medido con el robot parado y los sectores del chasis medidos; el renderizado es esquemático "
                      "(no fotorrealista) y la cámara simulada no reproduce la oscuridad ni el desenfoque reales (brillo 65/255, nitidez < 60). "
                      "No se afirma que la simulación sea indistinguible de la imagen real.")

    # ---------------------------------------------------------------- 7b guardián
    doc.add_heading("7b. Guardián del sendero y patas de mesas/sillas", 1)
    doc.add_paragraph("Escenarios nuevos con patas (círculos de 3 y 1,7 cm de radio + respaldos): lab_patas_mesas, aula_sillas, comedor_mesas_sillas, bajo_mesa_larga, sillas_sueltas. "
                      "guardian_sendero.py mide el error contra la referencia, corrige con pulsos, correlaciona LIDAR + IMU + cámara + error con una memoria de casos "
                      "(qué se hizo en situaciones parecidas y cuánto mejoró) y APUNTA cada corrección real (JSONL + memoria). Apagado en el robot (guardian_sendero_activo=false).")
    gr = js(RAIZ / "docs" / "guardian" / "resultados.json")
    if gr:
        modos = [("base", "base (explorador de hoy)"), ("ley", "guardián, solo ley"), ("memoria", "guardián + memoria entrenada")]
        filas = []
        for m, nom in modos:
            l = [r for r in gr if r["modo"] == m]
            filas.append([nom, f"{100 * sum(r['exito'] for r in l) / len(l):.1f}", f"{sum(r['contactos'] for r in l) / len(l):.2f}",
                          f"{sum(r['rmse_cm'] for r in l) / len(l):.1f}", f"{sum(r['pico_cm'] for r in l) / len(l):.1f}", f"{sum(r['progreso'] for r in l) / len(l):.2f}"])
        tabla(doc, filas, ["modo", "éxito %", "contactos/ep", "RMSE cm", "pico cm", "progreso"])
        escs = sorted({(r["ent"], r["esc"]) for r in gr})
        fig, ax = plt.subplots(figsize=(11, 4.4))
        x = np.arange(len(escs))
        for i, (m, nom) in enumerate(modos):
            ax.bar(x + (i - 1) * 0.27, [np.mean([r["rmse_cm"] for r in gr if r["modo"] == m and (r["ent"], r["esc"]) == k]) for k in escs], 0.27, label=nom)
        ax.set_xticks(x); ax.set_xticklabels([k[1] for k in escs], rotation=65, ha="right", fontsize=7); ax.set_ylabel("RMSE al sendero (cm)"); ax.set_yscale("log"); ax.legend(fontsize=7)
        fig.tight_layout(); fig.savefig(FIG / "guardian_rmse.png", dpi=130); plt.close(fig)
        figura(doc, FIG / "guardian_rmse.png", 16, "Error medio contra el sendero por escenario (escala log; 3 semillas nuevas por escenario).")
        doc.add_paragraph("Lectura honesta: el grueso de la mejora viene de la ley de control (Stanley con el giro por pulso medido); la memoria por correlación de sensores "
                          "es neutra-positiva en simulación (la física simulada es consistente, así que no tiene mucho que aprender). Su valor esperado está en el robot real, "
                          "donde el suelo, el agarre y la deriva cambian: eso NO está demostrado. En la cámara del simulador el sendero se deriva del error verdadero con ruido (sintética). "
                          "No mejora: curva_pasillo (empeora levemente), callejón de 1 m, eslalon, laberintos ni pasillos con esquina (esos dependen de las maniobras man_*).")
    figura(doc, RAIZ / "docs" / "mapas_interiores.png", 16, "Mapas con las patas de mesas y sillas (fila inferior).")
    mc = js(RAIZ / "docs" / "digitalizacion" / "mapa_scan_ancla.json")
    if mc:
        doc.add_paragraph("Ancla de rumbo por paredes en MapaScan (banco con deriva 1-3 %): precisión 94-97 % sin ella y no mejora con ella; se deja apagada.")

    # ---------------------------------------------------------------- 8 contradicciones
    doc.add_heading("8. Contradicciones y correcciones", 1)
    for t in [
        "MANIFEST_1.md §0 dice que el signo invertido de la velocidad de rf2o está «sin corregir», pero la memoria del proyecto lo da por corregido el 2026-10-23: no hay entrada en el Historial que lo respalde. A verificar en el robot.",
        "entrenar_intensivo.py: el espacio de búsqueda no incluía las maniobras (man_*) que el simulador sí tiene; los mapas con esquinas no se podían resolver. Añadidas (8 parámetros).",
        "entorno_calibrado.py declaraba una «sala en L» pendiente porque el robot simulado atraviesa la pared; el simulador no bloquea físicamente las paredes (solo cuenta contactos). Sigue siendo así: un contacto en una esquina cuenta como fallo pero no detiene el modelo.",
        "git: tools/kaggle/publicar/dataset/{dataset,real}.npz aparecen borrados en el árbol de trabajo y no se incluyen en el commit de este informe.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("9. No verificado / pendiente", 1)
    for t in ["Nada de lo entrenado se probó en el robot real; las maniobras man_* no están portadas al explorador.",
              "El visor 3D interior (static/interior_render.js) no se pudo abrir en un navegador desde esta sesión: sin prueba visual.",
              "El detector de pulgón no está validado con fotos etiquetadas de las coles del operador.",
              "Los datos ambientales del SCD40 siguen siendo simulados (ver memoria)."]:
        doc.add_paragraph(t, style="List Bullet")
    OUT.mkdir(parents=True, exist_ok=True)
    ruta = OUT / "REPORTE_MEJORAS_20261024.docx"
    doc.save(ruta)
    print("Guardado", ruta)


if __name__ == "__main__":
    main()
